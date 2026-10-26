"""
sheet_db_manager_sp.py
======================
Direct SharePoint sheet-backed persistence for the DSM-NIS automation pipeline.

- Uses Microsoft Graph Workbook APIs directly for reads/writes.
- No workbook download to local disk in SharePoint mode.
- Auto-creates SharePoint workbook if missing.
"""

from __future__ import annotations

import json
import threading
import time
import urllib.parse
from datetime import datetime
from pathlib import Path
from typing import Optional

import requests
from openpyxl import Workbook, load_workbook

import config
import sharepoint_upload as sp

# ── Step-name constants ──────────────────────────────────────────────────────
STEP_READ_EMAIL           = "read_email"
STEP_EMAIL_EXTRACTION     = "email_extraction"
STEP_EMAIL_CLASSIFICATION = "email_classification"
STEP_EMAIL_VALIDATION     = "email_validation"
STEP_READ_VENDOR          = "read_vendor_master"
STEP_UPDATE_TRACKER       = "update_excel_tracker"
STEP_NIS_BOOKING          = "nis_booking"


def step_nis(field: str) -> str:
    return f"nis_booking_{field}"


def step_sap(field: str) -> str:
    return f"sap_automation_{field}"


# ── Status constants ─────────────────────────────────────────────────────────
STATUS_PENDING = "pending"
STATUS_RUNNING = "running"
STATUS_DONE    = "done"
STATUS_FAILED  = "failed"
STATUS_SKIPPED = "skipped"

_FIXED_STEPS = [
    STEP_READ_EMAIL,
    STEP_EMAIL_EXTRACTION,
    STEP_EMAIL_CLASSIFICATION,
    STEP_EMAIL_VALIDATION,
    STEP_READ_VENDOR,
    STEP_UPDATE_TRACKER,
]

_EMAIL_RUNS_HEADERS = [
    "id", "message_id", "subject", "received_at", "spv_name", "status",
    "body_html", "metadata_json", "spv_dsm_json", "vendor_data_json",
    "positive_fields_json", "negative_fields_json", "downloaded_pdf_names",
    "nis_checklist_id", "nis_checklist_pdf_uploaded", "checklist_value",
    "sap_checklist_pdf_uploaded",
    "booking_folder_path", "failure_email_sent", "nis_completion_email_sent_at",
    "created_at", "started_at", "finished_at", "error_message",
]

_RUN_STEPS_HEADERS = [
    "id", "run_id", "step_name", "status", "detail", "started_at",
    "finished_at", "error", "retry_count",
]

_NIS_COMPLETION_RECIPIENTS_SHEET = "nis_completion_recipients"
_NIS_COMPLETION_RECIPIENTS_HEADERS = ["emails"]

_DEFAULT_RUN_VALUES = {
    "spv_name": "", "status": STATUS_PENDING, "body_html": "", "metadata_json": "",
    "spv_dsm_json": "", "vendor_data_json": "", "positive_fields_json": "",
    "negative_fields_json": "", "downloaded_pdf_names": "",
    "nis_checklist_id": "", "nis_checklist_pdf_uploaded": False,
    "checklist_value": "", "sap_checklist_pdf_uploaded": False,
    "booking_folder_path": "", "failure_email_sent": 0, "nis_completion_email_sent_at": "",
    "started_at": "", "finished_at": "", "error_message": "",
}

_DEFAULT_STEP_VALUES = {
    "status": STATUS_PENDING, "detail": "", "started_at": "",
    "finished_at": "", "error": "", "retry_count": 0,
}

_LOCK = threading.RLock()
_DB_PATH = Path(getattr(config, "SHEET_DB_PATH", "automation_sheet_db.xlsx"))
_SP_ENABLED = bool(getattr(config, "SHEET_DB_USE_SHAREPOINT", True))
_SP_FOLDER = str(getattr(config, "SHEET_DB_SP_FOLDER", getattr(config, "SP_BOOKING_BASE", "")) or "").strip("/")
_SP_FILE_NAME = str(getattr(config, "SHEET_DB_FILE_NAME", "automation_sheet_db.xlsx") or "automation_sheet_db.xlsx")
_ITEM_ID_CACHE = ""


def _sp_file_path() -> str:
    return f"{_SP_FOLDER}/{_SP_FILE_NAME}".strip("/") if _SP_FOLDER else _SP_FILE_NAME


def _now() -> str:
    return datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ")


def _encode_sp_path(path: str) -> str:
    return "/".join(urllib.parse.quote(seg, safe="") for seg in path.split("/") if seg)


def _col_letter(col_idx: int) -> str:
    out = ""
    n = col_idx
    while n > 0:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


def _create_empty_workbook_bytes() -> bytes:
    from io import BytesIO

    wb = Workbook()
    ws_runs = wb.active
    ws_runs.title = "email_runs"
    ws_runs.append(_EMAIL_RUNS_HEADERS)
    ws_steps = wb.create_sheet("run_steps")
    ws_steps.append(_RUN_STEPS_HEADERS)
    ws_recipients = wb.create_sheet(_NIS_COMPLETION_RECIPIENTS_SHEET)
    ws_recipients.append(_NIS_COMPLETION_RECIPIENTS_HEADERS)

    bio = BytesIO()
    wb.save(bio)
    wb.close()
    return bio.getvalue()


def _local_ensure_workbook_exists():
    if _DB_PATH.exists():
        return
    _DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_runs = wb.active
    ws_runs.title = "email_runs"
    ws_runs.append(_EMAIL_RUNS_HEADERS)
    ws_steps = wb.create_sheet("run_steps")
    ws_steps.append(_RUN_STEPS_HEADERS)
    ws_recipients = wb.create_sheet(_NIS_COMPLETION_RECIPIENTS_SHEET)
    ws_recipients.append(_NIS_COMPLETION_RECIPIENTS_HEADERS)
    wb.save(_DB_PATH)
    wb.close()


def _graph_headers(content_type: bool = False) -> dict:
    h = dict(sp._headers())
    if content_type:
        h["Content-Type"] = "application/json"
    return h


def _graph_drive_url() -> str:
    return f"https://graph.microsoft.com/v1.0/drives/{sp._get_drive_id()}"


def _sp_item_id() -> str:
    global _ITEM_ID_CACHE
    if _ITEM_ID_CACHE:
        return _ITEM_ID_CACHE

    if _SP_FOLDER:
        sp.ensure_sp_folder(_SP_FOLDER)

    url = f"{_graph_drive_url()}/root:/{_encode_sp_path(_sp_file_path())}"
    resp = requests.get(url, headers=_graph_headers(), timeout=60)
    if resp.status_code == 200:
        _ITEM_ID_CACHE = resp.json()["id"]
        return _ITEM_ID_CACHE

    if resp.status_code != 404:
        raise RuntimeError(f"Could not resolve SharePoint sheet DB file: {resp.status_code} {resp.text[:200]}")

    uploaded = sp.upload_bytes_to_sp(_create_empty_workbook_bytes(), _SP_FOLDER, _SP_FILE_NAME)
    if not uploaded:
        raise RuntimeError(f"Could not create SharePoint sheet DB workbook at {_sp_file_path()}")

    resp2 = requests.get(url, headers=_graph_headers(), timeout=60)
    resp2.raise_for_status()
    _ITEM_ID_CACHE = resp2.json()["id"]
    print(f"[SHEET-DB] Created SharePoint workbook: {_sp_file_path()}")
    return _ITEM_ID_CACHE


_TRANSIENT_STATUS_CODES = {429, 502, 503, 504}
_MAX_RETRIES = 4
_BASE_BACKOFF = 5  # seconds


def _is_transient(exc_or_resp) -> bool:
    """Return True for errors that are worth retrying."""
    if isinstance(exc_or_resp, (requests.exceptions.ConnectionError,
                                requests.exceptions.Timeout)):
        return True
    if hasattr(exc_or_resp, "status_code"):
        return exc_or_resp.status_code in _TRANSIENT_STATUS_CODES
    return False


def _is_unsupported_workbook_response(resp) -> bool:
    body = (getattr(resp, "text", "") or "").lower()
    return (
        resp.status_code in (400, 409, 500, 501)
        and ("unsupportedworkbook" in body or "filecorrupttryrepair" in body)
    )


def _sp_recreate_workbook(reason: str):
    global _ITEM_ID_CACHE
    print(f"[SHEET-DB] Recreating SharePoint workbook after Excel rejected it: {reason[:160]}")
    uploaded = sp.upload_bytes_to_sp(_create_empty_workbook_bytes(), _SP_FOLDER, _SP_FILE_NAME)
    if not uploaded:
        raise RuntimeError(f"Could not recreate SharePoint sheet DB workbook at {_sp_file_path()}")
    _ITEM_ID_CACHE = ""
    time.sleep(3)
    _sp_item_id()


def _sp_ensure_sheet(sheet_name: str, retry_on_recreate: bool = True):
    item_id = _sp_item_id()
    url = f"{_graph_drive_url()}/items/{item_id}/workbook/worksheets/add"
    last_exc = None
    for attempt in range(_MAX_RETRIES + 1):
        try:
            resp = requests.post(url, headers=_graph_headers(content_type=True),
                                 json={"name": sheet_name}, timeout=60)
        except (requests.exceptions.ConnectionError, requests.exceptions.Timeout) as exc:
            last_exc = exc
            if attempt < _MAX_RETRIES:
                wait = _BASE_BACKOFF * (2 ** attempt)
                print(f"[SHEET-DB] Transient error ensuring '{sheet_name}' (attempt {attempt+1}): {exc} → retry in {wait}s")
                time.sleep(wait)
                continue
            raise RuntimeError(f"Failed ensuring worksheet {sheet_name} after {_MAX_RETRIES} retries: {exc}") from exc

        if resp.status_code in (200, 201, 400, 409):
            return
        if retry_on_recreate and _is_unsupported_workbook_response(resp):
            _sp_recreate_workbook(resp.text)
            return _sp_ensure_sheet(sheet_name, retry_on_recreate=False)
        if _is_transient(resp) and attempt < _MAX_RETRIES:
            wait = _BASE_BACKOFF * (2 ** attempt)
            print(f"[SHEET-DB] Transient {resp.status_code} ensuring '{sheet_name}' (attempt {attempt+1}) → retry in {wait}s")
            time.sleep(wait)
            continue
        raise RuntimeError(f"Failed ensuring worksheet {sheet_name}: {resp.status_code} {resp.text[:200]}")
    raise RuntimeError(f"Failed ensuring worksheet {sheet_name} after {_MAX_RETRIES} retries (last: {last_exc})")


def _sp_get_values(sheet_name: str) -> list[list]:
    _sp_ensure_sheet(sheet_name)
    item_id = _sp_item_id()
    url = f"{_graph_drive_url()}/items/{item_id}/workbook/worksheets('{sheet_name}')/usedRange(valuesOnly=true)?$select=values"
    last_exc = None
    for attempt in range(_MAX_RETRIES + 1):
        try:
            resp = requests.get(url, headers=_graph_headers(), timeout=60)
        except (requests.exceptions.ConnectionError, requests.exceptions.Timeout) as exc:
            last_exc = exc
            if attempt < _MAX_RETRIES:
                wait = _BASE_BACKOFF * (2 ** attempt)
                print(f"[SHEET-DB] Transient error reading '{sheet_name}' (attempt {attempt+1}): {exc} → retry in {wait}s")
                time.sleep(wait)
                continue
            raise RuntimeError(f"Failed reading worksheet {sheet_name} after {_MAX_RETRIES} retries: {exc}") from exc

        if resp.status_code == 200:
            return resp.json().get("values", []) or []
        if resp.status_code == 404:
            return []
        if _is_transient(resp) and attempt < _MAX_RETRIES:
            wait = _BASE_BACKOFF * (2 ** attempt)
            print(f"[SHEET-DB] Transient {resp.status_code} reading '{sheet_name}' (attempt {attempt+1}) → retry in {wait}s")
            time.sleep(wait)
            continue
        raise RuntimeError(f"Failed reading worksheet {sheet_name}: {resp.status_code} {resp.text[:200]}")
    raise RuntimeError(f"Failed reading worksheet {sheet_name} after {_MAX_RETRIES} retries (last: {last_exc})")


def _sp_write_values(sheet_name: str, values: list[list]):
    _sp_ensure_sheet(sheet_name)
    if not values:
        return

    item_id = _sp_item_id()
    rows = len(values)
    cols = max((len(r) for r in values), default=1)
    end = f"{_col_letter(cols)}{rows}"

    clear_url = f"{_graph_drive_url()}/items/{item_id}/workbook/worksheets('{sheet_name}')/range(address='A1:AZ20000')/clear"
    try:
        requests.post(clear_url, headers=_graph_headers(content_type=True), json={"applyTo": "All"}, timeout=60)
    except Exception:
        pass

    write_url = f"{_graph_drive_url()}/items/{item_id}/workbook/worksheets('{sheet_name}')/range(address='A1:{end}')"
    last_exc = None
    for attempt in range(_MAX_RETRIES + 1):
        try:
            resp = requests.patch(write_url, headers=_graph_headers(content_type=True),
                                  json={"values": values}, timeout=60)
        except (requests.exceptions.ConnectionError, requests.exceptions.Timeout) as exc:
            last_exc = exc
            if attempt < _MAX_RETRIES:
                wait = _BASE_BACKOFF * (2 ** attempt)
                print(f"[SHEET-DB] Transient error writing '{sheet_name}' (attempt {attempt+1}): {exc} → retry in {wait}s")
                time.sleep(wait)
                continue
            raise RuntimeError(f"Failed writing worksheet {sheet_name} after {_MAX_RETRIES} retries: {exc}") from exc

        if resp.status_code in (200, 201):
            return
        if _is_transient(resp) and attempt < _MAX_RETRIES:
            wait = _BASE_BACKOFF * (2 ** attempt)
            print(f"[SHEET-DB] Transient {resp.status_code} writing '{sheet_name}' (attempt {attempt+1}) → retry in {wait}s")
            time.sleep(wait)
            continue
        raise RuntimeError(f"Failed writing worksheet {sheet_name}: {resp.status_code} {resp.text[:200]}")
    raise RuntimeError(f"Failed writing worksheet {sheet_name} after {_MAX_RETRIES} retries (last: {last_exc})")


def _local_get_values(sheet_name: str) -> list[list]:
    _local_ensure_workbook_exists()
    wb = load_workbook(_DB_PATH)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return []
    ws = wb[sheet_name]
    out = []
    for row in ws.iter_rows(values_only=True):
        out.append(["" if c is None else c for c in row])
    wb.close()
    return out


def _local_write_values(sheet_name: str, values: list[list]):
    _local_ensure_workbook_exists()
    wb = load_workbook(_DB_PATH)
    if sheet_name in wb.sheetnames:
        idx = wb.sheetnames.index(sheet_name)
        del wb[sheet_name]
        ws = wb.create_sheet(sheet_name, idx)
    else:
        ws = wb.create_sheet(sheet_name)
    for r in values:
        ws.append(r)
    wb.save(_DB_PATH)
    wb.close()


def _backend_get_values(sheet_name: str) -> list[list]:
    return _sp_get_values(sheet_name) if _SP_ENABLED else _local_get_values(sheet_name)


def _backend_write_values(sheet_name: str, values: list[list]):
    if _SP_ENABLED:
        _sp_write_values(sheet_name, values)
    else:
        _local_write_values(sheet_name, values)


def _sp_write_range(sheet_name: str, start_row: int, values: list[list]):
    _sp_ensure_sheet(sheet_name)
    if not values:
        return

    item_id = _sp_item_id()
    rows = len(values)
    cols = max((len(r) for r in values), default=1)
    end_col = _col_letter(cols)
    end_row = start_row + rows - 1

    write_url = f"{_graph_drive_url()}/items/{item_id}/workbook/worksheets('{sheet_name}')/range(address='A{start_row}:{end_col}{end_row}')"
    last_exc = None
    for attempt in range(_MAX_RETRIES + 1):
        try:
            resp = requests.patch(write_url, headers=_graph_headers(content_type=True),
                                  json={"values": values}, timeout=60)
        except (requests.exceptions.ConnectionError, requests.exceptions.Timeout) as exc:
            last_exc = exc
            if attempt < _MAX_RETRIES:
                wait = _BASE_BACKOFF * (2 ** attempt)
                time.sleep(wait)
                continue
            raise RuntimeError(f"Failed writing range to {sheet_name} after {_MAX_RETRIES} retries: {exc}") from exc

        if resp.status_code in (200, 201):
            return
        if _is_transient(resp) and attempt < _MAX_RETRIES:
            wait = _BASE_BACKOFF * (2 ** attempt)
            time.sleep(wait)
            continue
        raise RuntimeError(f"Failed writing range to {sheet_name}: {resp.status_code} {resp.text[:200]}")
    raise RuntimeError(f"Failed writing range to {sheet_name} after {_MAX_RETRIES} retries (last: {last_exc})")


def _local_write_range(sheet_name: str, start_row: int, values: list[list]):
    _local_ensure_workbook_exists()
    wb = load_workbook(_DB_PATH)
    if sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
    else:
        ws = wb[sheet_name]
    
    for i, row in enumerate(values):
        for j, val in enumerate(row):
            ws.cell(row=start_row + i, column=j + 1, value=val)
            
    wb.save(_DB_PATH)
    wb.close()


def _backend_write_range(sheet_name: str, start_row: int, values: list[list]):
    if _SP_ENABLED:
        _sp_write_range(sheet_name, start_row, values)
    else:
        _local_write_range(sheet_name, start_row, values)


def _rows_from_values(values: list[list], expected_headers: list[str]) -> list[dict]:
    if not values:
        return []
    header = [str(v or "").strip() for v in (values[0] if values else [])]
    if not any(header):
        header = list(expected_headers)
    out = []
    for raw in values[1:]:
        row = {h: (raw[i] if i < len(raw) else "") for i, h in enumerate(header)}
        if any((str(v).strip() if v is not None else "") for v in row.values()):
            out.append(row)
    return out


def _values_from_rows(headers: list[str], rows: list[dict]) -> list[list]:
    return [list(headers)] + [[r.get(h, "") for h in headers] for r in rows]


def _ensure_sheet_headers(sheet_name: str, headers: list[str]):
    values = _backend_get_values(sheet_name)
    if values and any(str(v or "").strip() for v in values[0]):
        return
    _backend_write_values(sheet_name, [list(headers)])


def _load_state() -> dict:
    runs = _rows_from_values(_backend_get_values("email_runs"), _EMAIL_RUNS_HEADERS)
    steps = _rows_from_values(_backend_get_values("run_steps"), _RUN_STEPS_HEADERS)
    for r in runs:
        r.setdefault("failure_email_sent", 0)
        r.setdefault("nis_completion_email_sent_at", "")
        r.setdefault("nis_checklist_pdf_uploaded", False)
        r.setdefault("sap_checklist_pdf_uploaded", False)
    for s in steps:
        s.setdefault("retry_count", 0)
    return {"email_runs": runs, "run_steps": steps}


def _migrate_email_runs_headers():
    values = _backend_get_values("email_runs")
    if not values:
        _backend_write_values("email_runs", [list(_EMAIL_RUNS_HEADERS)])
        return

    current_headers = [str(v or "").strip() for v in values[0]]
    if current_headers == _EMAIL_RUNS_HEADERS:
        return

    rows = _rows_from_values(values, _EMAIL_RUNS_HEADERS)
    _backend_write_values("email_runs", _values_from_rows(_EMAIL_RUNS_HEADERS, rows))


def _save_state(state: dict):
    _backend_write_values("email_runs", _values_from_rows(_EMAIL_RUNS_HEADERS, state["email_runs"]))
    _backend_write_values("run_steps", _values_from_rows(_RUN_STEPS_HEADERS, state["run_steps"]))


def _next_id(rows: list[dict]) -> int:
    m = 0
    for r in rows:
        try:
            m = max(m, int(r.get("id") or 0))
        except Exception:
            continue
    return m + 1


def _find_row(rows: list[dict], **where) -> Optional[dict]:
    for r in rows:
        ok = True
        for k, v in where.items():
            if str(r.get(k, "")) != str(v):
                ok = False
                break
        if ok:
            return r
    return None


def _sort_by_created(rows: list[dict]) -> list[dict]:
    return sorted(rows, key=lambda r: str(r.get("created_at") or ""))


# ─────────────────────────────────────────────────────────────────────────────
# Public API (drop-in for main.py)
# ─────────────────────────────────────────────────────────────────────────────
def init_db():
    with _LOCK:
        if _SP_ENABLED:
            _sp_item_id()
        _ensure_sheet_headers(_NIS_COMPLETION_RECIPIENTS_SHEET, _NIS_COMPLETION_RECIPIENTS_HEADERS)
        _migrate_email_runs_headers()
        state = _load_state()
        if not state["email_runs"] and not state["run_steps"]:
            _save_state({"email_runs": [], "run_steps": []})
    print(f"[SHEET-DB] Initialised → {_sp_file_path() if _SP_ENABLED else _DB_PATH}")


def is_email_registered(message_id: str) -> bool:
    return get_run_by_message_id(message_id) is not None


def register_email(message_id: str, subject: str, received_at: str, body_html: str = "") -> str:
    now = _now()
    with _LOCK:
        runs_vals = _backend_get_values("email_runs")
        steps_vals = _backend_get_values("run_steps")
        
        runs_headers = [str(v or "").strip() for v in (runs_vals[0] if runs_vals else _EMAIL_RUNS_HEADERS)]
        msg_idx = runs_headers.index("message_id") if "message_id" in runs_headers else -1
        
        for i, row in enumerate(runs_vals):
            if i == 0: continue
            if msg_idx >= 0 and msg_idx < len(row) and str(row[msg_idx]) == str(message_id):
                id_idx = runs_headers.index("id") if "id" in runs_headers else -1
                return str(row[id_idx]) if id_idx >= 0 and id_idx < len(row) else ""
                
        run_id = str(message_id or "")
        
        new_run_dict = {
            **_DEFAULT_RUN_VALUES,
            "id": run_id,
            "message_id": message_id,
            "subject": subject,
            "received_at": received_at,
            "body_html": body_html,
            "status": STATUS_PENDING,
            "created_at": now,
        }
        
        new_run_row = [new_run_dict.get(h, "") for h in runs_headers]
        _backend_write_range("email_runs", len(runs_vals) + 1, [new_run_row])
        
        steps_headers = [str(v or "").strip() for v in (steps_vals[0] if steps_vals else _RUN_STEPS_HEADERS)]
        
        id_idx = steps_headers.index("id") if "id" in steps_headers else 0
        max_id = 0
        for i, row in enumerate(steps_vals):
            if i == 0: continue
            if id_idx < len(row):
                try: max_id = max(max_id, int(row[id_idx] or 0))
                except: pass
                
        new_step_rows = []
        for step in _FIXED_STEPS:
            max_id += 1
            sd = {
                **_DEFAULT_STEP_VALUES,
                "id": max_id,
                "run_id": run_id,
                "step_name": step,
                "status": STATUS_PENDING,
            }
            new_step_rows.append([sd.get(h, "") for h in steps_headers])
            
        _backend_write_range("run_steps", len(steps_vals) + 1, new_step_rows)
        
    print(f"[SHEET-DB] Registered  run_id={run_id[:40]}  msg={message_id[:40]}…")
    return run_id


def ensure_dynamic_step(run_id: str, step_name: str):
    with _LOCK:
        steps_vals = _backend_get_values("run_steps")
        if not steps_vals: return
        headers = [str(v or "").strip() for v in steps_vals[0]]
        
        run_idx = headers.index("run_id") if "run_id" in headers else -1
        name_idx = headers.index("step_name") if "step_name" in headers else -1
        
        for i, row in enumerate(steps_vals):
            if i == 0: continue
            r_run = str(row[run_idx]) if run_idx >= 0 and run_idx < len(row) else ""
            r_name = str(row[name_idx]) if name_idx >= 0 and name_idx < len(row) else ""
            if r_run == str(run_id) and r_name == str(step_name):
                return
                
        id_idx = headers.index("id") if "id" in headers else 0
        max_id = 0
        for i, row in enumerate(steps_vals):
            if i == 0: continue
            if id_idx < len(row):
                try: max_id = max(max_id, int(row[id_idx] or 0))
                except: pass
                
        sd = {
            **_DEFAULT_STEP_VALUES,
            "id": max_id + 1,
            "run_id": str(run_id),
            "step_name": step_name,
            "status": STATUS_PENDING,
        }
        new_row = [sd.get(h, "") for h in headers]
        _backend_write_range("run_steps", len(steps_vals) + 1, [new_row])


def get_run_by_message_id(message_id: str) -> Optional[dict]:
    with _LOCK:
        return _find_row(_load_state()["email_runs"], message_id=message_id)


def get_run_by_id(run_id: str) -> Optional[dict]:
    with _LOCK:
        rows = _load_state()["email_runs"]
    for r in rows:
        if str(r.get("id") or "") == str(run_id) or str(r.get("message_id") or "") == str(run_id):
            return r
    return None


def get_pending_runs() -> list:
    with _LOCK:
        rows = _load_state()["email_runs"]
    return _sort_by_created([r for r in rows if r.get("status") in (STATUS_PENDING, STATUS_RUNNING)])


def get_incomplete_runs() -> list:
    with _LOCK:
        rows = _load_state()["email_runs"]
    return _sort_by_created([r for r in rows if r.get("status") in (STATUS_PENDING, STATUS_RUNNING, STATUS_FAILED)])


def _update_run(run_id: str, fields: dict):
    with _LOCK:
        runs_vals = _backend_get_values("email_runs")
        if not runs_vals:
            raise ValueError(f"email_runs is empty")
        headers = [str(v or "").strip() for v in runs_vals[0]]
        
        id_idx = headers.index("id") if "id" in headers else -1
        msg_idx = headers.index("message_id") if "message_id" in headers else -1
        
        row_idx = -1
        for i, row in enumerate(runs_vals):
            if i == 0: continue
            r_id = str(row[id_idx]) if id_idx >= 0 and id_idx < len(row) else ""
            r_msg = str(row[msg_idx]) if msg_idx >= 0 and msg_idx < len(row) else ""
            if r_id == str(run_id) or r_msg == str(run_id):
                row_idx = i
                break
                
        if row_idx == -1:
            raise ValueError(f"run_id={run_id} not found")
            
        row_data = list(runs_vals[row_idx])
        while len(row_data) < len(headers):
            row_data.append("")
            
        for k, v in fields.items():
            if k in headers:
                row_data[headers.index(k)] = v
                
        _backend_write_range("email_runs", row_idx + 1, [row_data])


def save_extracted_data(run_id: str, metadata_row: dict, spv_dsm_rows: list):
    _update_run(run_id, {"metadata_json": json.dumps(metadata_row or {}), "spv_dsm_json": json.dumps(spv_dsm_rows or [])})


def save_classification(run_id: str, positive: dict, negative: dict):
    _update_run(run_id, {"positive_fields_json": json.dumps(positive or {}), "negative_fields_json": json.dumps(negative or {})})


def save_vendor_data(run_id: str, vendor_data: dict):
    _update_run(run_id, {"vendor_data_json": json.dumps(vendor_data or {})})


def save_downloaded_pdfs(run_id: str, pdf_paths: list):
    save_downloaded_pdf_name(run_id, "dsm_pdf", pdf_paths)


def _parse_downloaded_pdf_names(raw) -> dict:
    if isinstance(raw, dict):
        return dict(raw)
    try:
        parsed = json.loads(raw or "{}")
    except Exception:
        return {}
    if isinstance(parsed, dict):
        return parsed
    if isinstance(parsed, list):
        return {"dsm_pdf": parsed}
    return {}


def save_downloaded_pdf_name(run_id: str, pdf_type: str, pdf_names):
    names = [
        str(p).replace("\\", "/").split("/")[-1]
        for p in (pdf_names if isinstance(pdf_names, (list, tuple, set)) else [pdf_names])
        if str(p or "").strip()
    ]
    with _LOCK:
        row = get_run_by_id(run_id) or {}
        name_map = _parse_downloaded_pdf_names(row.get("downloaded_pdf_names"))
        name_map[str(pdf_type)] = names if len(names) != 1 else names[0]
        _update_run(run_id, {"downloaded_pdf_names": json.dumps(name_map)})


def _parse_id_collection(raw) -> list[str]:
    raw_text = str(raw or "").strip()
    if not raw_text:
        return []
    try:
        parsed = json.loads(raw_text)
    except Exception:
        parsed = None
    if isinstance(parsed, list):
        return [str(v).strip() for v in parsed if str(v or "").strip()]
    if isinstance(parsed, dict):
        return [str(v).strip() for v in parsed.values() if str(v or "").strip()]
    return [part.strip() for part in raw_text.split(",") if part.strip()]


def save_nis_checklist_id(run_id: str, checklist_id: str, append: bool = False):
    checklist_id = str(checklist_id or "").strip()
    if not append:
        _update_run(run_id, {"nis_checklist_id": checklist_id})
        return
    row = get_run_by_id(run_id) or {}
    ids = _parse_id_collection(row.get("nis_checklist_id"))
    if checklist_id and checklist_id not in ids:
        ids.append(checklist_id)
    _update_run(run_id, {"nis_checklist_id": ", ".join(ids)})


def save_checklist_value(run_id: str, checklist_value: str):
    _update_run(run_id, {"checklist_value": str(checklist_value or "")})


def save_booking_folder_path(run_id: str, folder_path: str):
    _update_run(run_id, {"booking_folder_path": str(folder_path or "")})


def get_run_by_nis_checklist_id(checklist_id: str) -> Optional[dict]:
    checklist_id = str(checklist_id or "").strip()
    with _LOCK:
        rows = _load_state()["email_runs"]
    for row in rows:
        if checklist_id in _parse_id_collection(row.get("nis_checklist_id")):
            return row
    return None


def get_run_by_sap_checklist_value(checklist_value: str) -> Optional[dict]:
    with _LOCK:
        return _find_row(_load_state()["email_runs"], checklist_value=checklist_value)

def mark_nis_checklist_pdf_uploaded(run_id: str, uploaded: bool = True):
    _update_run(run_id, {"nis_checklist_pdf_uploaded": bool(uploaded)})


def mark_sap_checklist_pdf_uploaded(run_id: str, uploaded: bool = True):
    _update_run(run_id, {"sap_checklist_pdf_uploaded": bool(uploaded)})


def save_spv_name(run_id: str, spv_name: str):
    _update_run(run_id, {"spv_name": spv_name})


def is_failure_email_sent(run_id: str) -> bool:
    row = get_run_by_id(run_id)
    return bool(row and int(row.get("failure_email_sent") or 0))


def mark_failure_email_sent(run_id: str):
    _update_run(run_id, {"failure_email_sent": 1})


def get_nis_completion_recipients() -> list[str]:
    with _LOCK:
        values = _backend_get_values(_NIS_COMPLETION_RECIPIENTS_SHEET)
    rows = _rows_from_values(values, _NIS_COMPLETION_RECIPIENTS_HEADERS)
    seen, recipients = set(), []
    for row in rows:
        raw = str(row.get("emails") or "").strip()
        if not raw:
            continue
        for addr in raw.replace(";", ",").split(","):
            email = addr.strip()
            key = email.lower()
            if email and key not in seen:
                seen.add(key)
                recipients.append(email)
    return recipients


def is_nis_completion_email_sent(run_id: str) -> bool:
    row = get_run_by_id(run_id)
    return bool(row and str(row.get("nis_completion_email_sent_at") or "").strip())


def mark_nis_completion_email_sent(run_id: str, sent_at: str = ""):
    _update_run(run_id, {"nis_completion_email_sent_at": sent_at or _now()})


def get_runs_pending_nis_completion_email() -> list[dict]:
    with _LOCK:
        rows = _load_state()["email_runs"]
    return _sort_by_created([
        r for r in rows
        if not str(r.get("nis_completion_email_sent_at") or "").strip()
    ])


def mark_run_started(run_id: str):
    _update_run(run_id, {"status": STATUS_RUNNING, "started_at": _now()})


def mark_run_done(run_id: str):
    _update_run(run_id, {"status": STATUS_DONE, "finished_at": _now()})


def mark_run_failed(run_id: str, error: str):
    _update_run(run_id, {"status": STATUS_FAILED, "finished_at": _now(), "error_message": str(error)[:2000]})


def get_steps(run_id: str) -> list:
    with _LOCK:
        steps = _load_state()["run_steps"]
    out = []
    for r in steps:
        if str(r.get("run_id") or "") == str(run_id):
            out.append(r)
    return sorted(out, key=lambda x: int(x.get("id") or 0))


def get_step(run_id: str, step_name: str) -> Optional[dict]:
    for row in get_steps(run_id):
        if str(row.get("step_name") or "") == str(step_name):
            return row
    return None


def _update_step(run_id: str, step_name: str, fields: dict):
    with _LOCK:
        steps_vals = _backend_get_values("run_steps")
        if not steps_vals:
            return
        headers = [str(v or "").strip() for v in steps_vals[0]]
        
        run_idx = headers.index("run_id") if "run_id" in headers else -1
        name_idx = headers.index("step_name") if "step_name" in headers else -1
        
        row_idx = -1
        for i, row in enumerate(steps_vals):
            if i == 0: continue
            r_run = str(row[run_idx]) if run_idx >= 0 and run_idx < len(row) else ""
            r_name = str(row[name_idx]) if name_idx >= 0 and name_idx < len(row) else ""
            if r_run == str(run_id) and r_name == str(step_name):
                row_idx = i
                break
                
        if row_idx == -1:
            id_idx = headers.index("id") if "id" in headers else 0
            max_id = 0
            for i, row in enumerate(steps_vals):
                if i == 0: continue
                if id_idx < len(row):
                    try: max_id = max(max_id, int(row[id_idx] or 0))
                    except: pass
            
            sd = {**_DEFAULT_STEP_VALUES, "id": max_id + 1, "run_id": str(run_id), "step_name": step_name}
            sd.update(fields)
            new_row = [sd.get(h, "") for h in headers]
            _backend_write_range("run_steps", len(steps_vals) + 1, [new_row])
        else:
            row_data = list(steps_vals[row_idx])
            while len(row_data) < len(headers):
                row_data.append("")
            for k, v in fields.items():
                if k in headers:
                    row_data[headers.index(k)] = v
            _backend_write_range("run_steps", row_idx + 1, [row_data])


def is_step_done(run_id: str, step_name: str) -> bool:
    s = get_step(run_id, step_name)
    return s is not None and s.get("status") == STATUS_DONE


def is_step_skipped(run_id: str, step_name: str) -> bool:
    s = get_step(run_id, step_name)
    return s is not None and s.get("status") == STATUS_SKIPPED


def are_all_sap_steps_complete(run_id: str) -> bool:
    rows = [r for r in get_steps(run_id) if str(r.get("step_name") or "").startswith("sap_automation_")]
    if not rows:
        return True
    return all(r.get("status") in (STATUS_DONE, STATUS_SKIPPED) for r in rows)


def reset_crashed_sap_steps(run_id: str, max_retries: int) -> int:
    reset_count = 0
    with _LOCK:
        steps_vals = _backend_get_values("run_steps")
        if not steps_vals: return 0
        headers = [str(v or "").strip() for v in steps_vals[0]]
        
        run_idx = headers.index("run_id") if "run_id" in headers else -1
        name_idx = headers.index("step_name") if "step_name" in headers else -1
        status_idx = headers.index("status") if "status" in headers else -1
        retry_idx = headers.index("retry_count") if "retry_count" in headers else -1
        err_idx = headers.index("error") if "error" in headers else -1
        fin_idx = headers.index("finished_at") if "finished_at" in headers else -1
        
        for i, row in enumerate(steps_vals):
            if i == 0: continue
            
            r_run = str(row[run_idx]) if run_idx >= 0 and run_idx < len(row) else ""
            if r_run != str(run_id): continue
                
            r_name = str(row[name_idx]) if name_idx >= 0 and name_idx < len(row) else ""
            if not r_name.startswith("sap_automation_"): continue
                
            r_status = str(row[status_idx]) if status_idx >= 0 and status_idx < len(row) else ""
            r_retry = int(row[retry_idx]) if retry_idx >= 0 and retry_idx < len(row) and row[retry_idx] else 0
            
            if r_status in (STATUS_RUNNING, STATUS_FAILED) and r_retry < max_retries:
                row_data = list(row)
                while len(row_data) < len(headers): row_data.append("")
                row_data[status_idx] = STATUS_PENDING
                if err_idx >= 0: row_data[err_idx] = ""
                if fin_idx >= 0: row_data[fin_idx] = ""
                _backend_write_range("run_steps", i + 1, [row_data])
                reset_count += 1
                
    return reset_count


def step_start(run_id: str, step_name: str, detail: str = ""):
    _update_step(run_id, step_name, {"status": STATUS_RUNNING, "started_at": _now(), "detail": detail, "error": "", "finished_at": ""})


def step_done(run_id: str, step_name: str, detail: str = ""):
    _update_step(run_id, step_name, {"status": STATUS_DONE, "finished_at": _now(), "detail": detail})


def step_failed(run_id: str, step_name: str, error: str):
    _update_step(run_id, step_name, {"status": STATUS_FAILED, "finished_at": _now(), "error": str(error)[:2000]})


def step_skip(run_id: str, step_name: str, reason: str = ""):
    _update_step(run_id, step_name, {"status": STATUS_SKIPPED, "finished_at": _now(), "detail": reason})


def get_step_retry_count(run_id: str, step_name: str) -> int:
    row = get_step(run_id, step_name)
    return int(row.get("retry_count") or 0) if row else 0


def increment_step_retry(run_id: int, step_name: str, error: str) -> int:
    row = get_step(run_id, step_name)
    current = int(row.get("retry_count") or 0) if row else 0
    new_count = current + 1
    _update_step(run_id, step_name, {"retry_count": new_count, "status": STATUS_PENDING, "error": str(error)[:2000], "finished_at": _now()})
    return new_count


def print_run_summary(run_id: int):
    run = get_run_by_id(run_id)
    if not run:
        print(f"\n[SHEET-DB] Run {run_id} not found\n")
        return
    steps = get_steps(run_id)
    icons = {STATUS_DONE: "✅", STATUS_FAILED: "❌", STATUS_RUNNING: "🔄", STATUS_SKIPPED: "⏭ ", STATUS_PENDING: "⏳"}
    print(f"\n{'═' * 68}")
    print(f"  RUN #{run_id}  |  {run.get('subject', '')}")
    print(f"  Overall  : {run.get('status', '')}")
    print(f"  SPV      : {run.get('spv_name', '')}")
    print(f"  Started  : {run.get('started_at', '')}   Finished : {run.get('finished_at', '')}")
    if run.get("error_message"):
        print(f"  Error    : {str(run.get('error_message'))[:120]}")
    print("  Steps    :")

    for s in steps:
        icon = icons.get(s.get("status"), "?")
        dur = ""
        if s.get("started_at") and s.get("finished_at"):
            try:
                fmt = "%Y-%m-%dT%H:%M:%SZ"
                secs = int((datetime.strptime(str(s.get("finished_at")), fmt) - datetime.strptime(str(s.get("started_at")), fmt)).total_seconds())
                dur = f" ({secs}s)"
            except Exception:
                pass
        retry_count = int(s.get("retry_count") or 0)
        retry_tag = f"  [retries={retry_count}]" if retry_count else ""
        print(f"    {icon}  {str(s.get('step_name', '')):<38} {s.get('status', '')}{dur}{retry_tag}")
        if s.get("error"):
            print(f"         └─ {str(s.get('error'))[:90]}")
    print(f"{'═' * 68}\n")
