"""
Build the Tech Assurance Group (TAG/ARB) Review document for the
DSM-NIS Booking Automation project, following the exact structure and styling
of the reference template (Tech Assurance Group Review Template_V1 1.docx).
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "assets", "adani_logo.jpeg")
DIAG_LOGICAL = os.path.join(HERE, "diagram_logical.png")
DIAG_LAYERED = os.path.join(HERE, "diagram_layered.png")
OUT = os.path.join(HERE, "DSM_NIS_Booking_TAG_ARB_Review.docx")

TEAL = RGBColor(0x0F, 0x5B, 0x52)
HDR_BG = "0F5B52"
GREY_BG = "D9D9D9"
LIGHT = "F2F2F2"

doc = Document()

# base style
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10)

# --- Page setup: narrower margins so wide tables have room ---
for section in doc.sections:
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)

# Usable content width (Letter 8.5in - 0.6 - 0.6 = 7.3in)
USABLE_W = 7.3


def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def set_cell_margins(cell, top=40, bottom=40, left=90, right=90):
    """Cell padding in twips (1/20 pt). Keeps text off the borders."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("start", left),
                     ("end", right), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)


def set_col_width(cell, width_in):
    """Force a cell's width via tcW so fixed layout is honored."""
    cell.width = Inches(width_in)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(width_in * 1440)))
    tcW.set(qn("w:type"), "dxa")


def fix_table_layout(table, col_widths):
    """Apply fixed layout + explicit grid widths + padding + vertical centering.
    col_widths: list of column widths in inches (should sum ~= USABLE_W)."""
    table.autofit = False
    table.allow_autofit = False
    # fixed layout so Word honors the widths and doesn't reflow
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    # total table width
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(int(sum(col_widths) * 1440)))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    # define the grid
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gc, w in zip(grid.findall(qn("w:gridCol")), col_widths):
            gc.set(qn("w:w"), str(int(w * 1440)))
    for row in table.rows:
        for cell, w in zip(row.cells, col_widths):
            set_col_width(cell, w)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_text(cell, text, bold=False, color=None, size=10, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def header_row(table, texts, size=10):
    row = table.rows[0]
    for i, t in enumerate(texts):
        shade_cell(row.cells[i], HDR_BG)
        set_cell_text(row.cells[i], t, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
                      size=size, align=WD_ALIGN_PARAGRAPH.CENTER)


def add_kv_table(rows, col0_w=2.5, col1_w=None, shade_keys=True):
    if col1_w is None:
        col1_w = USABLE_W - col0_w
    t = doc.add_table(rows=len(rows), cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(rows):
        if shade_keys:
            shade_cell(t.rows[i].cells[0], LIGHT)
        set_cell_text(t.rows[i].cells[0], k, bold=True)
        set_cell_text(t.rows[i].cells[1], v)
    fix_table_layout(t, [col0_w, col1_w])
    return t


def h1(text):
    p = doc.add_heading(level=1)
    run = p.add_run(text)
    run.font.color.rgb = TEAL
    return p


def para(text, italic=False, size=10):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(size)
    return p


# =====================================================================
# COVER / HEADER
# =====================================================================
if os.path.exists(LOGO):
    lp = doc.add_paragraph()
    lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    lp.add_run().add_picture(LOGO, width=Inches(1.3))

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Tech Assurance Group Review Template")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = TEAL

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run("DSM-NIS Booking Automation")
rs.bold = True
rs.font.size = Pt(14)

doc.add_paragraph()

# Revision History
p = doc.add_paragraph()
p.add_run("Revision History:").bold = True
rev = doc.add_table(rows=2, cols=6)
rev.style = "Table Grid"
header_row(rev, ["No.", "Version", "Date of Publishing", "Author", "Reviewer", "Approver"])
vals = ["1", "V1.0", "10th September 2026", "EA Team",
        "Vikas Baweja, Gopal Chaubey, Sagar Gulati, Rajesh Lad, Manoj Gupta, Biswajit Baruah, Kallol Sarkar",
        ""]
for i, v in enumerate(vals):
    set_cell_text(rev.rows[1].cells[i], v, size=9)
fix_table_layout(rev, [0.45, 0.7, 1.3, 0.8, 2.65, 1.4])

doc.add_paragraph()
tp = doc.add_paragraph()
tr = tp.add_run("Tech Assurance Group Reviews (To be proportionately filled)")
tr.bold = True
tr.font.size = Pt(14)
tr.font.color.rgb = TEAL
para("This is the single TAG submission template. All projects start with the Impact Assessment "
     "form. Based on the assessment outcome, the system classifies the request into NR (No Review), "
     "LR (Light Review), or FR (Full Review). Only the sections marked applicable for the assigned "
     "category must be completed.", italic=True)

# =====================================================================
# 1. BASIC DETAILS
# =====================================================================
h1("1. Basic Details (applicable to all categories)")
add_kv_table([
    ("Project Name", "DSM-NIS Booking Automation (AGEL142)"),
    ("Business Unit", "Adani Green Energy Limited (AGEL)"),
    ("Sub-BU / Location", "Renewables – Finance / Commercial Operations"),
    ("Functional Owner", "Harshit Shah (Business User)"),
    ("BU IT Owner", "Kiran KR"),
    ("Application/Project Type", "AI / Automation (Digital project)"),
    ("Business Analyst (BA)", "Priyanka Pandey"),
    ("Developer", "Rohit Kumar (AI Engineer)"),
    ("Project Timeline",
     "Initiated: 09-Feb-2026  |  Phase 1 completed: 11-Apr-2026  |  "
     "Phase 2 started: 17-Jun-2026 (in progress). Development complete; deployment pending."),
    ("High-level Objective",
     "Automate the weekly DSM (Deviation Settlement Mechanism) / NIS (Non-Invoice Settlement) "
     "booking process end-to-end: read DSM emails, extract SPV charge tables, book positive amounts "
     "on the NIS Fiori portal and negative amounts in SAP, track HOD/FC approvals, merge evidence "
     "PDFs, update the SharePoint tracker, and send stage-wise completion notifications."),
])

# =====================================================================
# 2. STANDARDIZATION & REUSE
# =====================================================================
h1("2. Standardization & Reuse Assessment (NR / LR / FR)")
para("Mandatory section enforcing the group standardization mandate.", italic=True)
add_kv_table([
    ("Capability", "Email-triggered RPA + AI booking automation across NIS portal and SAP GUI"),
    ("Similar solution existing in the Group. (provide reference to an existing solution or prior "
     "TAG/ARB ID)", "No"),
    ("Overlap Estimate (Significant / Moderate / No overlap)", "No overlap"),
    ("Seeking Exception? (Yes/No)", "No"),
    ("Exit / Convergence Strategy (in case of exception)", "Not applicable"),
    ("Delta scope that needs to be reviewed (In case of reuse)", "Not applicable"),
])

# =====================================================================
# 3. BUSINESS REQUIREMENTS
# =====================================================================
h1("3. Business Requirements (NR / LR / FR)")
para("Approved high-level business requirements for the DSM-NIS Booking Automation.", italic=True)
reqs = [
    ("Monitor the dedicated Outlook mailbox and detect DSM-NIS-BOOKING emails from the last 2 days.", "Must-Have"),
    ("Extract metadata and SPV charge tables from the email HTML body and download invoice attachments.", "Must-Have"),
    ("Classify each SPV amount: positive -> NIS booking, negative -> SAP checklist, zero -> skip.", "Must-Have"),
    ("Look up vendor master data (vendor/company code, cost center, plant, bank key) from SharePoint.", "Must-Have"),
    ("Automate NIS Fiori portal booking for positive amounts (login, PDF upload, expense items, submit).", "Must-Have"),
    ("Automate SAP ZDCC debit/credit checklist creation for negative amounts and capture checklist number.", "Must-Have"),
    ("Run booking on a schedule (3x daily) with crash-resume and idempotent re-processing.", "Must-Have"),
    ("Group same-cost-center rows into one NIS booking; support 1/2/3/5-row table scenarios.", "Must-Have"),
    ("Extract invoice numbers from vendor PDFs using AI (Azure OpenAI), with auto-generation fallback.", "Must-Have"),
    ("Track HOD and FC approvals via inbound emails and progress the workflow automatically.", "Must-Have"),
    ("Merge all evidence documents (email, invoice, NIS & SAP checklists) into one PDF per booking.", "Must-Have"),
    ("Trigger the Debit/Credit finance email flow for 5-row tables and process finance reply PDFs.", "Must-Have"),
    ("Update the SharePoint output tracker (DSM-NIS-OUTPUT.xlsx) per SPV with all booking references.", "Must-Have"),
    ("Send stage-wise and final completion notification emails to the configured recipients.", "Must-Have"),
    ("Run NIS bookings in parallel (up to 3 concurrent) to reduce processing time.", "Nice-to-Have"),
    ("Send an LLM-generated failure-explanation email when a booking cannot be completed.", "Nice-to-Have"),
]
t = doc.add_table(rows=len(reqs) + 1, cols=3)
t.style = "Table Grid"
header_row(t, ["#", "High Level Business Requirements", "Must Have / Nice to Have"])
for i, (desc, pri) in enumerate(reqs, start=1):
    set_cell_text(t.rows[i].cells[0], str(i), align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t.rows[i].cells[1], desc)
    set_cell_text(t.rows[i].cells[2], pri, align=WD_ALIGN_PARAGRAPH.CENTER)
fix_table_layout(t, [0.4, 5.5, 1.4])

para("Process flow diagram (email trigger -> classification -> NIS/SAP booking -> approvals -> "
     "merge -> completion) is embedded in Section 5 below.")

# =====================================================================
# 4. VENDOR DETAILS
# =====================================================================
h1("4. Vendor Details (NR / LR / FR)")
para("4.a  No Review – Not applicable.")
para("4.b  Light Review – Not applicable.")
para("4.c  Vendor details:")
add_kv_table([
    ("Development Model", "In-house build (Adani AI Engineering)"),
    ("External Product Vendor", "None – no third-party product licence procured"),
    ("Underlying Platform Vendors",
     "Microsoft Azure (PostgreSQL, OpenAI, AD/Graph, Windows VM), SAP (SAP GUI), "
     "SharePoint / Microsoft 365"),
    ("Rationale", "Solution is a custom automation over existing licensed Adani platforms "
                  "(Microsoft 365, Azure, SAP, NIS). No new vendor onboarding required."),
])

# =====================================================================
# 5. SOLUTION OVERVIEW & ARCHITECTURE SNAPSHOT
# =====================================================================
h1("5. Solution Overview & Architecture Snapshot (LR / FR)")
add_kv_table([
    ("Deployment Model (SaaS / Adani Cloud / On Prem)",
     "Adani Cloud (Azure) + On-Prem/VM. NIS service runs on a host with Playwright; SAP service "
     "runs on a Windows VM with SAP GUI. Shared Azure PostgreSQL for state."),
    ("Core functional components/modules",
     "Email ingestion & classification (Graph API), HTML table parser, vendor master lookup, "
     "NIS booking engine (Playwright), SAP automation engine (win32com COM), AI invoice extractor, "
     "HOD/FC approval tracking, PDF merge, Debit/Credit finance flow, SharePoint tracker, "
     "notification service."),
    ("User Type (Internal/External)", "Internal"),
    ("Who will be the users?",
     "AGEL Finance / Commercial team (business user Harshit Shah), sender of DSM-NIS-BOOKING emails; "
     "HOD & FC approvers; Finance team (Debit/Credit notes)."),
    ("Internet Facing (Yes/No)", "No (runs within the Adani network / tenant; portals accessed as an "
                                 "authenticated internal user)"),
])

para("Logical Architecture – DSM-NIS Booking Automation:")
if os.path.exists(DIAG_LOGICAL):
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(DIAG_LOGICAL, width=Inches(7.0))

para("Layered Technical Architecture:")
if os.path.exists(DIAG_LAYERED):
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(DIAG_LAYERED, width=Inches(7.0))

# =====================================================================
# 6. TECHNOLOGY STACK
# =====================================================================
h1("6. Technology Stack (LR / FR)")
para("List of key technologies used. Deviations from group standards are marked explicitly "
     "(Standard vs Exception).", italic=True)
tech = [
    ("Automation Runtime", "Python 3.10+", "3.11", "Standard", "Windows VM / host process"),
    ("Browser Automation", "Playwright (Microsoft Edge / Chromium)", "1.44+", "Standard", "Windows VM / host"),
    ("SAP Automation", "win32com COM + pywinauto (SAP GUI Scripting)", "pywin32 306+", "Standard", "Windows VM (SAP GUI)"),
    ("Database", "Azure PostgreSQL (managed) / SQLite (local dev)", "psycopg2 2.9+", "Standard", "Azure – India region"),
    ("Email & SharePoint", "Microsoft Graph API via MSAL + requests", "msal 1.28+", "Standard", "Microsoft 365 tenant"),
    ("AI / LLM", "Azure OpenAI (GPT) – invoice extraction & failure explanation", "GPT-4 class", "Standard", "Azure OpenAI"),
    ("PDF Processing", "pdfplumber (extract), pypdf (merge), reportlab (email->PDF)", "—", "Standard", "Host process"),
    ("Job Scheduler", "Windows Task Scheduler (3x daily)", "—", "Standard", "Windows VM"),
    ("Version Control / CI-CD", "GitLab", "—", "Standard", "GitLab"),
]
t = doc.add_table(rows=len(tech) + 1, cols=5)
t.style = "Table Grid"
header_row(t, ["Layer / Component", "Technology / Product", "Version", "Standard / Exception",
               "Underlying Compute (where deployed?)"], size=9)
for i, row in enumerate(tech, start=1):
    for c, val in enumerate(row):
        set_cell_text(t.rows[i].cells[c], val, size=9)
fix_table_layout(t, [1.35, 2.05, 1.0, 1.15, 1.75])

# =====================================================================
# 7. INTEGRATION INTERFACE REGISTER
# =====================================================================
h1("7. Integration Interface Register (LR / FR)")
integ = [
    ("Outlook Mailbox (M365)", "DSM-NIS Automation", "Read DSM emails, approvals, D/C replies; send notifications",
     "Microsoft Graph REST API", "Point-to-Point", "Azure AD OAuth2 (client credentials, MSAL)",
     "Email body HTML tables, PDF attachments, notification emails", "Scheduled (3x daily) / On-Demand"),
    ("SharePoint / OneDrive (M365)", "DSM-NIS Automation", "Vendor master, output tracker, booking folder storage",
     "Microsoft Graph REST API", "Point-to-Point", "Azure AD OAuth2 (client credentials)",
     "Vendor master Excel, tracker Excel, booking PDFs (invoice/checklist/merged)", "On-Demand / per booking"),
    ("Azure OpenAI", "DSM-NIS Automation", "Invoice number extraction from PDFs; failure explanations",
     "REST API", "Point-to-Point", "API Key",
     "Extracted PDF text (prompt), invoice number / explanation (response)", "On-Demand"),
    ("NIS Fiori Launchpad", "DSM-NIS Automation", "Non-PO vendor payment booking of positive amounts",
     "Browser automation (Playwright/Edge)", "UI Automation", "User login + password + RSA/TOTP OTP",
     "Booking form fields, invoice PDF upload, NIS checklist number", "Per booking (parallel up to 3)"),
    ("SAP ERP (SAP GUI, ZDCC)", "DSM-NIS Automation", "Debit/Credit checklist creation for negative amounts",
     "SAP GUI Scripting (win32com COM)", "UI Automation", "SAP GUI session (user credentials)",
     "Debit/credit form fields, SAP checklist number, checklist PDF", "Per booking / on Windows VM"),
    ("Azure PostgreSQL", "DSM-NIS Automation", "Shared pipeline state & resume tracking",
     "psycopg2 (TCP, SSL)", "Point-to-Point", "DB username/password (SSL enforced)",
     "email_runs, booking_rows, run_steps, debit_credit_* records", "Real-time / per step"),
]
t = doc.add_table(rows=len(integ) + 1, cols=8)
t.style = "Table Grid"
header_row(t, ["From System", "To System", "Purpose", "Integration Method",
               "Middleware / Orchestration", "Authentication Method", "Data being exchanged",
               "Mode & Frequency"], size=7.5)
for i, row in enumerate(integ, start=1):
    for c, val in enumerate(row):
        set_cell_text(t.rows[i].cells[c], val, size=7.5)
# widths sum to 7.3; tighter cell padding for this dense table
fix_table_layout(t, [0.95, 0.85, 1.05, 0.95, 0.9, 0.95, 0.85, 0.8])
for row in t.rows:
    for cell in row.cells:
        set_cell_margins(cell, top=30, bottom=30, left=55, right=55)

# =====================================================================
# 8. NFRs
# =====================================================================
h1("8. Non-Functional Requirements (NFRs) (LR / FR)")
para("Measurable targets (initial assumptions where actuals are pending deployment validation).", italic=True)
add_kv_table([
    ("Deployment Location", "Azure Cloud – India region (state DB) + Windows VM for portal/SAP automation"),
    ("Availability Target (%)", "Batch automation – runs 3x daily. Azure PostgreSQL SLA 99.99%; "
                                "target ~99% successful scheduled-run completion."),
    ("RTO requirement", "24 Hours (re-run picks up incomplete work from DB state)"),
    ("RPO requirement", "Near-zero for pipeline state (every step persisted before/after execution)"),
    ("HA Strategy", "Stateless services + durable state in Azure PostgreSQL. Any host can resume from "
                    "the shared DB; SAP work is polled independently by the SAP service."),
    ("DR Strategy", "Recovery via Azure-managed PostgreSQL backups and redeployment of the VM/host. "
                    "Idempotent, resume-aware pipeline avoids duplicate bookings on recovery."),
    ("Backup Strategy", "Azure PostgreSQL managed automated backups; SharePoint/OneDrive versioning "
                        "for tracker and booking documents."),
    ("Sizing & Scalability", "Configurable parallel NIS booking (NIS_PARALLEL_COUNT, default 3). "
                             "Throughput scales with worker count and VM capacity."),
    ("Performance Requirements", "Single scan of last 2 days processed within one scheduled run; "
                                 "per-row NIS booking within a few minutes including login + OTP."),
    ("Concurrent Users / Transactions", "System (unattended) actor; up to 3 concurrent NIS booking "
                                        "sessions. No interactive concurrent end users."),
])

# =====================================================================
# 9. KEY RISK & DEVIATION SUMMARY
# =====================================================================
h1("9. Key Risk & Deviation Summary (LR / FR)")
risks = [
    ("Credential & Access Management",
     "Automation depends on service-account credentials, RSA/OTP for NIS, SAP GUI login, and Graph API "
     "secrets. Delays in obtaining cost-center access, ID creation, RSA bypass, email credentials and "
     "CEO approval extended the timeline and remain an operational dependency.",
     "Store secrets outside source control (env / .env / secret store), rotate regularly, use dedicated "
     "service accounts, and document access dependencies for renewals."),
    ("Portal / GUI UI Changes",
     "NIS Fiori portal or SAP ZDCC screen changes can break UI-based automation (Playwright selectors / "
     "SAP scripting).",
     "Resilient selectors, retry limits (NIS_MAX_ROW_RETRIES), failure alerts, and a maintenance runbook "
     "to update flows when screens change."),
    ("Email Parsing Accuracy",
     "Malformed HTML tables, missing columns, or unexpected date formats in the source email may cause "
     "extraction errors or wrong classification.",
     "Strict validation step, date normalization, mandatory-field checks, and an LLM-generated failure "
     "email describing the missing/mismatched value."),
    ("AI Invoice Extraction Reliability",
     "LLM may return an incorrect or empty invoice number from vendor PDFs.",
     "Deterministic fallbacks (auto-generated invoice numbers for special SPVs/vendors), per-row "
     "pre-extraction for 3-row/INOX bills, and human-visible tracker for review."),
    ("Duplicate / Double Booking",
     "Re-runs or concurrent workers could book the same transaction twice.",
     "Message-id dedup, step-level DB state, DB locks on tracker writes, cost-center grouping, and "
     "idempotent approval handling."),
    ("Security & Data Handling",
     "Financial data, vendor details, and portal credentials are processed; unauthorized access or "
     "leakage would be sensitive.",
     "Internal-only execution, OAuth2 client-credential access scoped to required resources, SSL to "
     "PostgreSQL, and least-privilege service accounts."),
    ("Deployment Dependency",
     "Development is complete but deployment is pending; production readiness (scheduling, monitoring, "
     "VM hardening) still to be finalized.",
     "Complete deployment checklist: Task Scheduler setup, monitoring/alerting, credential provisioning, "
     "and a production smoke test before go-live."),
]
t = doc.add_table(rows=len(risks) + 1, cols=3)
t.style = "Table Grid"
header_row(t, ["Area", "Risk / Exposure", "Mitigation / Control"])
for i, (a, rk, m) in enumerate(risks, start=1):
    set_cell_text(t.rows[i].cells[0], a, bold=True, size=9)
    set_cell_text(t.rows[i].cells[1], rk, size=9)
    set_cell_text(t.rows[i].cells[2], m, size=9)
fix_table_layout(t, [1.5, 2.9, 2.9])

# =====================================================================
# 10. COST & COMMERCIAL VIEW (TCO)
# =====================================================================
h1("10. Cost & Commercial View (TCO) (FR mandatory, LR optional)")
para("High-level TCO for transparency. Detailed procurement artifacts are not required in this "
     "template. (3-year TCO to be shared.)", italic=True)
cost = [
    ("Software License / Subscription", "Recurring", "Nil (net-new)",
     "Built on existing Adani M365, Azure and SAP licences; no new product licence."),
    ("Infra & Services Cost", "Recurring", "To be confirmed",
     "Azure PostgreSQL, Azure OpenAI usage, Windows VM compute."),
    ("AI / LLM Usage", "Recurring", "To be confirmed",
     "Azure OpenAI token consumption for invoice extraction & failure explanations."),
    ("Annual Maintenance Cost", "Recurring", "To be confirmed",
     "In-house support for UI-flow updates, monitoring and enhancements."),
    ("Development (one-time)", "One time", "Internal effort",
     "In-house AI Engineering (Phase 1 + Phase 2)."),
]
t = doc.add_table(rows=len(cost) + 1, cols=4)
t.style = "Table Grid"
header_row(t, ["Cost Component", "Nature of expense (One time/Recurring)", "Value", "Notes"])
for i, row in enumerate(cost, start=1):
    for c, val in enumerate(row):
        set_cell_text(t.rows[i].cells[c], val, size=9)
fix_table_layout(t, [1.9, 1.7, 1.3, 2.4])

# =====================================================================
# 11. FULL REVIEW - DEEP DIVE
# =====================================================================
h1("11. Full Review – Deep Dive (FR only)")
para("Required only for Full Review items:")
for line in [
    "Detailed architecture diagrams: see Section 5 (logical + layered). Two-service design "
    "(NIS Booking + SAP Automation) communicating only through a shared PostgreSQL state store.",
    "Capacity and scalability analysis: parallel NIS booking is configurable (default 3 concurrent "
    "sessions); throughput scales with VM capacity and worker count; SAP work polled independently.",
    "Threat model (lite): internal-only execution; OAuth2 client-credential scopes; SSL to DB; "
    "secrets kept out of source control; least-privilege service accounts; RSA/TOTP for NIS login.",
    "Operational readiness: step-level DB tracking with crash-resume, idempotent operations, retry "
    "limits, per-SPV file logs, failure notification emails, and SAP recovery markers.",
    "Alternatives considered: manual booking (baseline, error-prone/slow) vs. UI automation over "
    "existing portals (chosen, no API available for NIS/SAP booking screens).",
    "Exception rationale + convergence plan: no group-standard deviation; UI automation is used only "
    "where no supported API exists for the target screens.",
]:
    doc.add_paragraph(line, style="List Bullet")

# =====================================================================
# 12. TAG DECISION
# =====================================================================
h1("12. Tech Assurance Group Decision (NR / LR / FR)")
add_kv_table([
    ("Decision (Approved / Rejected / Approved with Condition)", ""),
    ("Key Recommendations", ""),
    ("Architecture Decision", ""),
])

doc.save(OUT)
print("ARB written to", OUT)
